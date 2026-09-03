"""loaders.py — versione completa (riferimento del docente).

Stesso nome e stessi import di `working/loaders.py`: cambia solo la cartella.
I TODO 1-5 sono il codice che il docente scrive a schermo dal vivo, con accanto
la frase da dire mentre lo scrive; i TODO 6-7 sono una delle soluzioni valide
dell'esercizio (per il CSV la granularità è una scelta, non c'è una risposta
giusta unica).
"""
import csv
import sys

from dataclasses import dataclass, field
from pathlib import Path

import docx

from bs4 import BeautifulSoup
from markdown_it import MarkdownIt
from pypdf import PdfReader

SCRIPTS = Path(__file__).parent.parent
sys.path.insert(0, str(SCRIPTS / "scaffolding"))

from ask import mostra, rispondi  # noqa: E402 — il banco di prova, già scritto

DATASET = SCRIPTS / "dataset"


# --------------------------------------------------------- il formato interno
@dataclass
class Document:
    """Quello che ogni loader deve restituire: il testo estratto e la sua provenienza.

    metadata tiene almeno `source` (da quale file arriva) e `type` (il formato);
    i loader che lavorano pagina per pagina aggiungono anche `page`.
    """
    text: str
    metadata: dict = field(default_factory=dict)


# ------------------------------------------------------- TODO 1 · col docente 🎤
def load_pdf(path):
    """Il PDF, una pagina per Document (pypdf).

    Da dire mentre si scrive: un PDF non è testo lineare, è un layout di blocchi
    posizionati. pypdf ci dà il testo di ogni pagina, uno alla volta — e visto
    che le pagine esistono, teniamo il numero nei metadata: servirà per citare
    la fonte e per capire da dove è uscito un risultato di ricerca.
    `extract_text()` può tornare None su una pagina senza layer di testo: da
    qui l'`or ""`.
    """
    reader = PdfReader(path)
    docs = []
    for numero, pagina in enumerate(reader.pages):
        testo = pagina.extract_text() or ""
        docs.append(Document(testo, {"source": str(path), "page": numero + 1, "type": "pdf"}))
    return docs


# ------------------------------------------------------- TODO 2 · col docente 🎤
def load_docx(path):
    """Il documento Word, dai paragrafi dell'XML (python-docx).

    Da dire mentre si scrive: un .docx è uno zip di XML — il testo c'è, ma è
    annegato nei tag. python-docx ce lo espone come lista di paragrafi, e noi li
    riuniamo. Notare che cambia solo il corpo della funzione: quello che esce è
    di nuovo un Document, quindi da qui in poi nessuno sa più che era Word.
    """
    documento = docx.Document(path)
    testo = "\n".join(p.text for p in documento.paragraphs if p.text)
    return [Document(testo, {"source": str(path), "type": "docx"})]


# ------------------------------------------------------- TODO 3 · col docente 🎤
NOISE = ["script", "style", "header", "nav", "aside", "footer"]


def load_html(path):
    """La pagina web, senza il boilerplate (BeautifulSoup).

    Da dire mentre si scrive — è il momento chiave della lezione: prima si
    scrive la versione SENZA il ciclo di decompose(), si lancia il main e si
    legge la risposta alla domanda sul telaio (esce «10 anni», che è sbagliata).
    Poi si aggiunge il ciclo, si rilancia e la stessa domanda dà «5 anni».
    Il contenuto sta in <main>; header, nav, aside e footer parlano dello stesso
    argomento con numeri diversi, e il modello si fida di quelli perché usano
    le parole della domanda. I numeri esatti sono in `conduzione-docente.md`.
    """
    with open(path, encoding="utf-8") as f:
        soup = BeautifulSoup(f, "html.parser")
    for tag in soup(NOISE):
        tag.decompose()
    testo = soup.get_text(separator="\n", strip=True)
    return [Document(testo, {"source": str(path), "type": "html"})]


# ------------------------------------------------------- TODO 4 · col docente 🎤
def load_markdown(path):
    """Il Markdown come testo, senza la sintassi (markdown-it-py).

    Da dire mentre si scrive: il formato più gentile, è già quasi testo. Ma i
    cancelletti, i trattini e i pipe della tabella non servono al modello e
    sporcano il prompt: lo passiamo dal parser (Markdown → HTML) e poi ne
    prendiamo il testo con lo stesso BeautifulSoup di un minuto fa.
    """
    sorgente = Path(path).read_text(encoding="utf-8")
    html = MarkdownIt().render(sorgente)
    testo = BeautifulSoup(html, "html.parser").get_text("\n", strip=True)
    return [Document(testo, {"source": str(path), "type": "markdown"})]


# ------------------------------------------------------- TODO 5 · col docente 🎤
def load(path):
    """La porta unica: apre qualunque formato conosciuto, guardando l'estensione.

    Da dire mentre si scrive: da qui in poi chi carica un documento non deve più
    sapere quale libreria serve per quale file — chiama load() e basta.
    Aggiungere un formato vuol dire aggiungere una riga qui e una funzione
    sopra: ed è esattamente quello che fate voi nell'esercizio.
    """
    ext = Path(path).suffix.lower()
    if ext == ".pdf":
        return load_pdf(path)
    if ext == ".docx":
        return load_docx(path)
    if ext in (".html", ".htm"):
        return load_html(path)
    if ext in (".md", ".markdown"):
        return load_markdown(path)
    if ext == ".txt":
        return load_txt(path)
    if ext == ".csv":
        return load_csv(path)
    raise ValueError(f"Formato non supportato: {ext}")


# --------------------------------------------------------------- TODO 6 · voi ⌨️
def load_txt(path):
    """Il file di testo semplice, in Document.

    Una delle soluzioni valide: il testo c'è già, serve solo leggerlo e
    metterlo nel contenitore con i metadata giusti.
    """
    testo = Path(path).read_text(encoding="utf-8")
    return [Document(testo, {"source": str(path), "type": "txt"})]


# --------------------------------------------------------------- TODO 7 · voi ⌨️
def load_csv(path):
    """Il CSV, in Document — qui una riga per Document.

    Una delle soluzioni valide, e la scelta da discutere al reveal: una riga =
    un Document rende ogni ordine un risultato di ricerca a sé (utile quando si
    cerca «l'ordine di Nordvent»), ma perde la vista d'insieme sul file. Chi ha
    fatto un solo Document con tutte le righe ha l'effetto opposto: risponde
    bene a «quanti ordini sono consegnati?» e male a «che cosa ha ordinato
    Nordvent?». Nessuna delle due è sbagliata: dipende dalle domande.
    """
    with open(path, encoding="utf-8", newline="") as f:
        righe = list(csv.DictReader(f))
    docs = []
    for numero, riga in enumerate(righe):
        testo = "\n".join(f"{campo}: {valore}" for campo, valore in riga.items())
        docs.append(Document(testo, {"source": str(path), "type": "csv", "row": numero + 1}))
    return docs


if __name__ == "__main__":
    DOMANDA_PDF = "Di quanto è cresciuto il fatturato nel 2025?"
    DOMANDA_DOCX = "Quale città è stata scelta per il secondo showroom, e con che budget?"
    DOMANDA_HTML = "Quanto dura la garanzia sul telaio della scrivania?"
    DOMANDA_MD = "Che cosa devo fare quando il display mostra E07?"
    DOMANDA_TXT = "Quante segnalazioni di codice E07 sono arrivate questa settimana?"
    DOMANDA_CSV = "Quanti ordini risultano consegnati?"

    print("### TODO 1 · PDF")
    try:
        docs = load_pdf(DATASET / "relazione-annuale.pdf")
        mostra(docs, "relazione-annuale.pdf")
        rispondi(docs, DOMANDA_PDF)
    except NotImplementedError:
        print("  TODO 1 ancora da fare")

    print("\n### TODO 2 · DOCX")
    try:
        docs = load_docx(DATASET / "verbale-consiglio.docx")
        mostra(docs, "verbale-consiglio.docx")
        rispondi(docs, DOMANDA_DOCX)
    except NotImplementedError:
        print("  TODO 2 ancora da fare")

    print("\n### TODO 3 · HTML")
    try:
        docs = load_html(DATASET / "faq-garanzia.html")
        mostra(docs, "faq-garanzia.html")
        rispondi(docs, DOMANDA_HTML)
    except NotImplementedError:
        print("  TODO 3 ancora da fare")

    print("\n### TODO 4 · Markdown")
    try:
        docs = load_markdown(DATASET / "guida-lumadesk.md")
        mostra(docs, "guida-lumadesk.md")
        rispondi(docs, DOMANDA_MD)
    except NotImplementedError:
        print("  TODO 4 ancora da fare")

    print("\n### TODO 5 · load(): la porta unica")
    for nome in ("relazione-annuale.pdf", "verbale-consiglio.docx",
                 "faq-garanzia.html", "guida-lumadesk.md",
                 "note-supporto.txt", "ordini.csv"):
        try:
            docs = load(DATASET / nome)
            print(f"  {nome:26} → {len(docs)} Document · {docs[0].metadata.get('type')}")
        except NotImplementedError:
            print(f"  {nome:26} → il loader di questo formato è ancora da fare")
        except ValueError as e:
            print(f"  {nome:26} → ValueError: {e}")

    print("\n### TODO 6 · TXT")
    try:
        docs = load_txt(DATASET / "note-supporto.txt")
        mostra(docs, "note-supporto.txt")
        rispondi(docs, DOMANDA_TXT)
    except NotImplementedError:
        print("  TODO 6 ancora da fare")

    print("\n### TODO 7 · CSV")
    try:
        docs = load_csv(DATASET / "ordini.csv")
        mostra(docs, "ordini.csv")
        rispondi(docs, DOMANDA_CSV)
    except NotImplementedError:
        print("  TODO 7 ancora da fare")
