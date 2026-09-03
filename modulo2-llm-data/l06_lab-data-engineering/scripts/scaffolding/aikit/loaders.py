"""loaders.py — i loader scritti a Modulo 2 · Lezione 4, promossi nel toolkit.

Non si tocca: si importa. Ogni formato ha la sua funzione, `load()` smista per
estensione, e tutte restituiscono la stessa cosa — una lista di `Document`.
È il motivo per cui da qui in poi nessuno deve più sapere da che file arriva
il testo.
"""
import csv
from dataclasses import dataclass, field
from pathlib import Path

import docx
from bs4 import BeautifulSoup
from markdown_it import MarkdownIt
from pypdf import PdfReader


@dataclass
class Document:
    """Il testo estratto e la sua provenienza.

    metadata tiene almeno `source` (da quale file arriva) e `type` (il formato);
    i loader che lavorano pagina per pagina aggiungono anche `page`.
    """
    text: str
    metadata: dict = field(default_factory=dict)


def load_pdf(path):
    """Il PDF, una pagina per Document (pypdf).

    `extract_text()` può tornare None su una pagina senza layer di testo: da
    qui l'`or ""`.
    """
    reader = PdfReader(path)
    docs = []
    for numero, pagina in enumerate(reader.pages):
        testo = pagina.extract_text() or ""
        docs.append(Document(testo, {"source": str(path), "page": numero + 1, "type": "pdf"}))
    return docs


def load_docx(path):
    """Il documento Word, dai paragrafi dell'XML (python-docx)."""
    documento = docx.Document(path)
    testo = "\n".join(p.text for p in documento.paragraphs if p.text)
    return [Document(testo, {"source": str(path), "type": "docx"})]


NOISE = ["script", "style", "header", "nav", "aside", "footer"]


def load_html(path):
    """La pagina web, senza il boilerplate (BeautifulSoup).

    I tag di NOISE si tolgono con decompose(): in una pagina web il contorno è
    marcato, e sei tag lo portano via. Quando NON è marcato — un piede di
    pagina dentro un PDF — non c'è niente da togliere per struttura, e serve
    un'euristica sul testo: è il lavoro di questa lezione.
    """
    with open(path, encoding="utf-8") as f:
        soup = BeautifulSoup(f, "html.parser")
    for tag in soup(NOISE):
        tag.decompose()
    testo = soup.get_text(separator="\n", strip=True)
    return [Document(testo, {"source": str(path), "type": "html"})]


def load_markdown(path):
    """Il Markdown come testo, senza la sintassi (markdown-it-py)."""
    sorgente = Path(path).read_text(encoding="utf-8")
    html = MarkdownIt().render(sorgente)
    testo = BeautifulSoup(html, "html.parser").get_text("\n", strip=True)
    return [Document(testo, {"source": str(path), "type": "markdown"})]


def load_txt(path):
    """Il file di testo semplice, in Document."""
    testo = Path(path).read_text(encoding="utf-8")
    return [Document(testo, {"source": str(path), "type": "txt"})]


def load_csv(path):
    """Il CSV, una riga per Document."""
    with open(path, encoding="utf-8", newline="") as f:
        righe = list(csv.DictReader(f))
    docs = []
    for numero, riga in enumerate(righe):
        testo = "\n".join(f"{campo}: {valore}" for campo, valore in riga.items())
        docs.append(Document(testo, {"source": str(path), "type": "csv", "row": numero + 1}))
    return docs


def load(path):
    """La porta unica: apre qualunque formato conosciuto, guardando l'estensione."""
    ext = Path(path).suffix.lower()
    if ext == ".pdf":
        return load_pdf(path)
    if ext == ".docx":
        return load_docx(path)
    if ext in (".html", ".htm"):
        return load_html(path)
    if ext in (".md", ".markdown"):
        return load_markdown(path)
    if ext == ".txt":
        return load_txt(path)
    if ext == ".csv":
        return load_csv(path)
    raise ValueError(f"Formato non supportato: {ext}")
